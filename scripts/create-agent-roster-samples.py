"""Generate fully synthetic workbooks in the existing roster import format.

Run with the managed roster Python. All names, appointments and qualifications
are authored test fixtures; no values are copied from real NGO workbooks.
"""

import json
import sys
from datetime import datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deployment" / "agent-samples"
OUT.mkdir(parents=True, exist_ok=True)
sys.path.insert(
    0, str(ROOT / "MCPcode/src/ngopilot_mcp/payloads/rostercopiilot/backend")
)

from app.importer import parse_division_workbook  # noqa: E402 - read-only vendored import
from app.services.master_data_bridge import bootstrap_master_data_from_division  # noqa: E402
from app.domain import MasterElder  # noqa: E402

book = Workbook()
sheet = book.active
sheet.title = "恆常服務"
sheet["A1"] = "合成示範團隊 — 不是真實 NGO 排班"
sheet["A2"] = "星期"
names = ["測試甲", "測試乙", "測試丙", "測試丁", "測試戊", "測試己"]
for column, name in enumerate(names, 3):
    sheet.cell(2, column, name).fill = PatternFill("solid", fgColor="FFFFFF00")
    sheet.column_dimensions[sheet.cell(2, column).column_letter].width = 24
for day, label in enumerate("一二三四五六"):
    first = 3 + day * 10
    sheet.merge_cells(start_row=first, end_row=first + 9, start_column=1, end_column=1)
    sheet.cell(first, 1, label)
    for offset, period in [(0, "上"), (5, "下")]:
        row = first + offset
        sheet.merge_cells(start_row=row, end_row=row + 4, start_column=2, end_column=2)
        sheet.cell(row, 2, period)
        sheet.cell(row, 4, "ESC")
sheet["C3"] = "HC:測試長者甲(EH)"
sheet["C4"] = "09:00-10:00 測試區"
for column in range(3, 9):
    sheet.cell(64, column, "08:30-17:30")
    sheet.cell(65, column, "A" if column % 2 else "B")
base = OUT / "roster-division-synthetic.xlsx"
book.save(base)

master = bootstrap_master_data_from_division(parse_division_workbook(base))
master.origin = "synthetic_demo_authored_2026-10-03"
for worker in master.workers:
    worker.gender = "F"
    worker.notes = "Synthetic fixture; no real employee."
    for fact in worker.skill_facts:
        fact.source = "manual"
        fact.evidence = "Authored qualification in synthetic demonstration fixture, not an NGO confirmation."
for elder in master.elders:
    elder.gender = "F"
    elder.district = "測試區"
master.elders.append(
    MasterElder(
        id="synthetic-elder-b",
        display_name="測試長者乙",
        aliases=["測試長者乙"],
        gender="F",
        district="測試區",
        notes="Synthetic fixture.",
    )
)
for value in master.rule_config.values.values():
    value.confirmed = True
    value.assumption = "Explicit rule of this synthetic fixture only."
(OUT / "roster-master-synthetic.json").write_text(
    master.model_dump_json(indent=2), encoding="utf-8"
)

book = Workbook()
sheet = book.active
sheet.title = "52026"  # Existing parser's worksheet contract, not the target week.
sheet["A1"] = "合成 HC 工作表 — 此週沒有新增 HC，固定服務由分工表提供"
for start in (1, 8, 15, 22, 29):
    for offset, label in enumerate(
        ["Case", "單位", "節數", "時間", "照顧員", "日期", "備註"]
    ):
        sheet.cell(3, start + offset, label)
book.save(OUT / "HC-synthetic.xlsx")

book = Workbook()
sheet = book.active
sheet.title = "1月"  # Existing parser's worksheet contract.
sheet["A1"] = "合成護送工作表 — 使用格內日期 2026-10-06"
for col, label in enumerate(
    [
        "日期",
        "星期",
        "時段",
        "姓名",
        "單位",
        "時間",
        "地點",
        "目的",
        "交通",
        "備註",
        "經手人",
        "填寫日",
    ],
    1,
):
    sheet.cell(5, col, label)
for col, value in enumerate(
    [
        datetime(2026, 10, 6),
        "二",
        "上午",
        "測試長者乙",
        "EH",
        time(10, 0),
        "測試中心",
        "合成示範",
        "步行",
        "",
        "測試員",
        datetime(2026, 10, 3),
    ],
    1,
):
    sheet.cell(6, col, value)
sheet["A6"].number_format = "yyyy-mm-dd"
sheet["F6"].number_format = "hh:mm"
book.save(OUT / "Escort-synthetic.xlsx")
print(
    json.dumps({"synthetic_files": [p.name for p in OUT.iterdir()]}, ensure_ascii=False)
)
