"""Use a supplied transcript, never a speech provider, in Luna-only mode."""

from pathlib import Path


def read_transcript(path: str) -> str:
    source = Path(path)
    if source.suffix.lower() == ".docx":
        from docx import Document

        document = Document(source)
        lines = [p.text for p in document.paragraphs]
        lines.extend(
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        text = "\n".join(lines)
    elif source.suffix.lower() in {".txt", ".md"}:
        text = source.read_text(encoding="utf-8-sig")
    else:
        raise ValueError("Luna 不支援錄音轉寫。請提供 TXT、Markdown 或 DOCX 文字稿。")
    if not text.strip():
        raise ValueError("文字稿沒有可讀內容。")
    if len(text) > 500_000:
        raise ValueError("文字稿太長，請分段處理（上限 500,000 字元）。")
    return text.strip()


def install_transcript_adapter(service, transcriber):
    transcriber.transcribe_audio = read_transcript
    service.transcribe_audio = read_transcript
